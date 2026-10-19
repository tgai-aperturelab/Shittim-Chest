"""Bounded document parsing worker. No command execution, network or office macros."""

from __future__ import annotations

import base64
import json
import sys
from collections.abc import Iterator
from io import BytesIO
from typing import Any, cast
from zipfile import ZipFile


def units(data: bytes, suffix: str) -> Iterator[dict[str, Any]]:
    if suffix in {".docx", ".xlsx"}:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(e.file_size for e in entries) > 128 * 1024 * 1024:
                raise ValueError("Expanded document exceeds parser limits")
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF requires a decrypted copy")
        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            yield {
                "kind": "page",
                "location": index + 1,
                "text": text,
                "requires_ocr": not bool(text.strip()),
            }
    elif suffix == ".docx":
        from docx import Document
        from docx.table import Table

        document = Document(BytesIO(data))
        for index, item in enumerate(document.iter_inner_content()):
            if isinstance(item, Table):
                for row_number, row in enumerate(item.rows):
                    yield {
                        "kind": "table_row",
                        "location": f"block:{index + 1}/row:{row_number + 1}",
                        "text": "\t".join(cell.text for cell in row.cells),
                    }
            else:
                yield {"kind": "paragraph", "location": index + 1, "text": item.text}
    elif suffix == ".xlsx":
        import openpyxl

        workbook = cast(Any, openpyxl.load_workbook)(
            BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            for sheet in workbook.worksheets:
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    yield {
                        "kind": "spreadsheet_row",
                        "location": f"{sheet.title}!{row_number}",
                        "text": "\t".join("" if value is None else str(value) for value in row),
                        "formulas_evaluated": False,
                    }
        finally:
            workbook.close()
    else:
        raise ValueError("Unsupported document format")


def parse_document(
    data: bytes, suffix: str, offset: int = 0, limit: int = 50, character_offset: int = 0
) -> dict[str, Any]:
    if offset < 0 or not 1 <= limit <= 500 or character_offset < 0:
        raise ValueError("Invalid document pagination")
    results: list[dict[str, Any]] = []
    remaining = 32000
    for index, unit in enumerate(units(data, suffix)):
        if index < offset:
            continue
        if len(results) >= limit or remaining <= 0:
            return {
                "units": results,
                "next_offset": index,
                "next_character_offset": 0,
                "truncated": True,
            }
        start = character_offset if index == offset else 0
        text = str(unit["text"])[start:]
        page = text[:remaining]
        results.append({**unit, "text": page})
        remaining -= len(page)
        if len(page) < len(text):
            return {
                "units": results,
                "next_offset": index,
                "next_character_offset": start + len(page),
                "truncated": True,
            }
    return {"units": results, "next_offset": None, "next_character_offset": 0, "truncated": False}


if __name__ == "__main__":
    try:
        request = json.loads(sys.stdin.buffer.read(90 * 1024 * 1024))
        value = parse_document(
            base64.b64decode(request["data"], validate=True),
            request["suffix"],
            request.get("offset", 0),
            request.get("limit", 50),
            request.get("character_offset", 0),
        )
        print(json.dumps({"parsed": True, **value}, ensure_ascii=True))
    except Exception as exc:
        print(json.dumps({"parsed": False, "error": type(exc).__name__}))
