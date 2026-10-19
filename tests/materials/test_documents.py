from io import BytesIO

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from shittim_core.materials.documents import read_document


async def test_word_body_order_and_complete_pagination():
    document = Document()
    document.add_paragraph("第一段")
    row = document.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text, row.cells[1].text = "名称", "数量"
    document.add_paragraph("最后一段")
    stream = BytesIO()
    document.save(stream)
    page = await read_document(stream.getvalue(), ".docx", limit=2)
    assert page["parsed"] and page["truncated"]
    assert [u["text"] for u in page["units"]] == ["第一段", "名称\t数量"]
    final = await read_document(stream.getvalue(), ".docx", offset=page["next_offset"])
    assert final["units"][0]["text"] == "最后一段" and not final["truncated"]


async def test_sheet_names_rows_and_formulas_are_preserved_not_evaluated():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "金额"
    sheet.append(["名字", "金额"])
    sheet.append(["甲", 12])
    sheet.append(["合计", "=SUM(B2:B2)"])
    stream = BytesIO()
    workbook.save(stream)
    result = await read_document(stream.getvalue(), ".xlsx")
    assert result["parsed"]
    assert result["units"][2]["location"] == "金额!3"
    assert "=SUM(B2:B2)" in result["units"][2]["text"]
    assert result["units"][2]["formulas_evaluated"] is False


async def test_pdf_extracts_real_text_and_marks_blank_page_without_ocr():
    writer = PdfWriter()
    page = writer.add_blank_page(300, 300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 20 200 Td (Evidence 42) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.add_blank_page(300, 300)
    stream = BytesIO()
    writer.write(stream)
    result = await read_document(stream.getvalue(), ".pdf")
    assert result["parsed"] and "Evidence 42" in result["units"][0]["text"]
    assert not result["units"][0]["requires_ocr"]
    assert result["units"][1]["requires_ocr"]


async def test_long_document_unit_can_be_read_without_losing_tail():
    document = Document()
    text = "中" * 33000 + "末尾"
    document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    first = await read_document(stream.getvalue(), ".docx")
    second = await read_document(
        stream.getvalue(),
        ".docx",
        offset=first["next_offset"],
        character_offset=first["next_character_offset"],
    )
    assert first["units"][0]["text"] + second["units"][0]["text"] == text


async def test_malformed_document_and_process_timeout_are_reported():
    invalid = await read_document(b"not a PDF", ".pdf")
    assert not invalid["parsed"] and invalid["error"]
    timeout = await read_document(b"data", ".pdf", timeout=0.001)
    assert timeout == {"parsed": False, "error": "parser_timeout"}
